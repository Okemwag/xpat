"""PDF, Word and Excel reports: same content, labels kept, numbers stay numbers, uploaded text never becomes a formula."""

import io
from decimal import Decimal
import pytest
from conftest import row
from floodcat.core.constants import TIERS
from floodcat.reporting.document import Table, build_document
from floodcat.reporting.formats import render
from floodcat.services.analysis import analyse

pytest.importorskip("fpdf")
pytest.importorskip("docx")
pytest.importorskip("openpyxl")


@pytest.fixture(scope="module")
def analysis(starter_rows):
    return analyse(starter_rows[:40])


DECISION = {
    "created_at": "2026-10-08 10:00",
    "decided_by_name": "Una Underwriter",
    "premium_100_kes": "5000000",
    "offered_share_pct": "20",
    "recommendation": {"outcome": "share", "recommended_share_pct": 12.5},
    "outcome": "share",
    "share_pct": "10",
    "overrode": True,
    "reason": "Broker could not confirm the coordinates.",
}
BRIEFING = {
    "headline": "Concrete drives the tail",
    "model": "fake",
    "prompt_version": "briefing-v1",
    "fact_count": 12,
    "unsupported_figures": [],
    "sections": [
        {"heading": "What the results say", "paragraphs": ["Loss rises with rarity."]}
    ],
    "checks": ["Check the depth assumption."],
}


def test_every_format_renders(analysis):
    doc = build_document(
        analysis,
        label="Test",
        organisation="Test Re",
        briefing=BRIEFING,
        decisions=[DECISION],
    )
    pdf, mime = render(doc, "pdf", analysis)
    assert pdf.startswith(b"%PDF") and mime == "application/pdf"
    for kind in ("docx", "xlsx"):
        data, _ = render(doc, kind, analysis)
        assert data.startswith(b"PK")


def test_pdf_states_labels_and_limitations(analysis):
    pypdf = pytest.importorskip("pypdf")
    doc = build_document(analysis, label="Test", decisions=[DECISION])
    text = " ".join(
        p.extract_text()
        for p in pypdf.PdfReader(io.BytesIO(render(doc, "pdf", analysis)[0])).pages
    )
    assert (
        "SYNTHETIC" in text
        and "PROXY" in text
        and "Limitations" in text
        and "Not a price" in text
    )
    assert "Take a smaller share" in text and "Underwriting decisions recorded" in text


def test_word_has_sections_tables_and_ai_label(analysis):
    from docx import Document

    doc = build_document(analysis, label="Test", briefing=BRIEFING)
    word = Document(io.BytesIO(render(doc, "docx", analysis)[0]))
    headings = [p.text for p in word.paragraphs if p.style.name.startswith("Heading")]
    assert "Headline figures" in headings and any(
        h.startswith("AI underwriting briefing") for h in headings
    )
    assert (
        len(word.tables) >= 6 and len(word.inline_shapes) == 2
    )  # EP curve and construction charts
    assert any("AI" in p.text for p in word.paragraphs)


def test_excel_keeps_numbers_synthetic_flag_and_totals(analysis):
    from openpyxl import load_workbook

    doc = build_document(analysis, label="Test")
    wb = load_workbook(io.BytesIO(render(doc, "xlsx", analysis)[0]))
    assert {"Summary", "Scenarios", "Construction", "Provenance", "Limitations"} <= set(
        wb.sheetnames
    )
    curve = {
        p["tier"]: Decimal(p["loss_kes"])
        for p in analysis["runs"]["baseline"]["ep_curve"]
    }
    for tier in TIERS:
        ws = wb[f"Props base {tier}"]
        header = [c.value for c in ws[3]]
        rows = list(ws.iter_rows(min_row=4, values_only=True))
        loss, synthetic = header.index("Loss"), header.index("Synthetic")
        assert len(rows) == analysis["modelled_count"]
        assert all(r[synthetic] == "true" for r in rows)  # AGENTS invariant 5
        assert all(isinstance(r[loss], (int, float)) for r in rows)
        assert abs(sum(Decimal(str(r[loss])) for r in rows) - curve[tier]) < Decimal(
            "0.1"
        ) * len(rows)  # invariant 3
    scen = wb["Scenarios"]
    assert scen.cell(4, 1).number_format == '"1-in-"#,##0' and isinstance(
        scen.cell(4, 4).value, (int, float)
    )


def test_excel_never_writes_uploaded_text_as_a_formula():
    from openpyxl import load_workbook

    rows = [row('=HYPERLINK("http://x","click")', source="+cmd|calc"), row("T-2")]
    report = analyse(rows)
    data, _ = render(build_document(report, label="=1+1"), "xlsx", report)
    wb = load_workbook(io.BytesIO(data))
    cells = [
        c
        for ws in wb.worksheets
        for r in ws.iter_rows()
        for c in r
        if isinstance(c.value, str) and c.value[:1] in "=+"
    ]
    assert cells and all(c.data_type == "s" for c in cells)


def test_appendix_tables_are_excel_only(analysis):
    doc = build_document(analysis, label="Test")
    appendix = [b for b in doc.blocks if isinstance(b, Table) and b.data_only]
    assert len(appendix) == len(TIERS) and all(
        t.sheet.startswith("Props") for t in appendix
    )


def test_ai_run_adds_comparison(starter_rows):
    from floodcat.ai.evidence import Evidence

    ev = Evidence(
        "ev-1",
        "test",
        "flooded",
        "",
        "CBD",
        -1.2864,
        36.8172,
        "manual",
        "drainage",
        0.9,
        True,
        approved=True,
        reviewer="R",
    )
    report = analyse(starter_rows[:40], evidence=[ev], ai_adjustment=True)
    doc = build_document(report, label="AI")
    assert any(isinstance(b, Table) and b.sheet == "AI evidence" for b in doc.blocks)
    assert sum(isinstance(b, Table) and b.data_only for b in doc.blocks) == 2 * len(
        TIERS
    )
